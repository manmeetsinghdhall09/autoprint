"""
AI File Sorter - sorts files into folders based on what is inside them,
using a local LLM running on your own GPU through Ollama (https://ollama.com).

Nothing leaves your PC: file contents are only sent to the Ollama server on
localhost. By default the script only PREVIEWS the plan; pass --apply to move
files, and --undo <log> to put everything back.

Examples:
    python sorter.py "C:/Users/me/Downloads"
    python sorter.py "C:/Users/me/Downloads" --dest "D:/Sorted" --apply
    python sorter.py "C:/Users/me/Downloads" --categories Invoices Taxes Photos Code School Other
    python sorter.py --undo sort_log_20261007_120000.json
"""
import argparse
import base64
import csv
import json
import os
import shutil
import sys
import time
import urllib.error
import urllib.request

DEFAULT_CATEGORIES = [
    "Finance", "Invoices & Receipts", "Work", "School", "Personal Documents",
    "Code", "Photos", "Screenshots", "Music", "Videos", "Books & Papers",
    "Installers", "Archives", "Other",
]

TEXT_EXTS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".log", ".ini",
    ".cfg", ".yaml", ".yml", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h",
    ".cs", ".go", ".rs", ".php", ".rb", ".sh", ".bat", ".ps1", ".sql", ".css",
    ".tex", ".rtf",
}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"}

# Files whose content a text model can't read: classified by extension instead.
EXTENSION_FALLBACK = {
    "Music": {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a", ".wma"},
    "Videos": {".mp4", ".mkv", ".avi", ".mov", ".wmv", ".webm"},
    "Installers": {".exe", ".msi", ".dmg", ".deb", ".apk"},
    "Archives": {".zip", ".rar", ".7z", ".tar", ".gz", ".iso"},
    "Photos": IMAGE_EXTS | {".heic", ".raw", ".cr2", ".nef"},
}

SKIP_DIRS = {"$recycle.bin", "system volume information", "windows",
             "program files", "program files (x86)", "appdata", ".git",
             "node_modules", "__pycache__"}

MAX_CHARS = 3000          # how much text from each file is shown to the model
MAX_IMAGE_BYTES = 8 * 1024 * 1024


# ---------------------------------------------------------------- extraction

def extract_text(path):
    """Return a short text snippet from the file, or '' if unreadable."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in TEXT_EXTS:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(MAX_CHARS)
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            text = ""
            for page in reader.pages[:3]:
                text += (page.extract_text() or "") + "\n"
                if len(text) >= MAX_CHARS:
                    break
            return text[:MAX_CHARS]
        if ext == ".docx":
            import docx  # python-docx
            text = "\n".join(p.text for p in docx.Document(path).paragraphs)
            return text[:MAX_CHARS]
        if ext in {".xlsx", ".xlsm"}:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets[:2]:
                for row in ws.iter_rows(max_row=30, values_only=True):
                    rows.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(rows)[:MAX_CHARS]
        if ext == ".pptx":
            from pptx import Presentation
            parts = []
            for slide in Presentation(path).slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return "\n".join(parts)[:MAX_CHARS]
    except ImportError as e:
        print(f"  (missing optional library for {ext}: {e.name} - run: pip install -r requirements.txt)")
    except Exception:
        pass
    return ""


def extension_guess(path, categories):
    ext = os.path.splitext(path)[1].lower()
    for cat, exts in EXTENSION_FALLBACK.items():
        if ext in exts and cat in categories:
            return cat
    return "Other" if "Other" in categories else categories[-1]


# ---------------------------------------------------------------- ollama

def ollama_chat(host, model, prompt, categories, image_path=None, timeout=300):
    schema = {
        "type": "object",
        "properties": {
            "category": {"type": "string", "enum": categories},
            "reason": {"type": "string"},
        },
        "required": ["category", "reason"],
    }
    message = {"role": "user", "content": prompt}
    if image_path:
        with open(image_path, "rb") as f:
            message["images"] = [base64.b64encode(f.read()).decode("ascii")]
    body = json.dumps({
        "model": model,
        "messages": [message],
        "format": schema,
        "stream": False,
        "options": {"temperature": 0},
    }).encode("utf-8")
    req = urllib.request.Request(f"{host}/api/chat", data=body,
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = json.loads(resp.read())
    result = json.loads(data["message"]["content"])
    if result.get("category") not in categories:
        raise ValueError(f"model returned unknown category {result.get('category')!r}")
    return result["category"], result.get("reason", "")


def check_ollama(host, models):
    try:
        with urllib.request.urlopen(f"{host}/api/tags", timeout=5) as resp:
            installed = {m["name"] for m in json.loads(resp.read())["models"]}
    except (urllib.error.URLError, OSError):
        sys.exit(f"Can't reach Ollama at {host}. Install it from https://ollama.com "
                 f"and make sure it is running.")
    for m in models:
        if m and m not in installed and f"{m}:latest" not in installed:
            sys.exit(f"Model '{m}' is not downloaded. Run:  ollama pull {m}")


def classify(path, args, categories):
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    cat_list = ", ".join(categories)

    if ext in IMAGE_EXTS and args.vision_model and os.path.getsize(path) <= MAX_IMAGE_BYTES:
        prompt = (f"Look at this image (file name: {name}) and pick the single best "
                  f"folder for it from: {cat_list}.")
        return ollama_chat(args.host, args.vision_model, prompt, categories, image_path=path)

    text = extract_text(path)
    if not text.strip():
        return extension_guess(path, categories), "classified by file type (no readable text)"

    prompt = (f"You are organizing a messy computer. Decide which folder this file "
              f"belongs in, based mainly on its CONTENT.\n"
              f"Folders: {cat_list}\n\n"
              f"File name: {name}\n"
              f"Content excerpt:\n\"\"\"\n{text}\n\"\"\"")
    return ollama_chat(args.host, args.model, prompt, categories)


# ---------------------------------------------------------------- planning

def iter_files(root, recursive, dest):
    dest = os.path.abspath(dest)
    for dirpath, dirnames, filenames in os.walk(root):
        # never descend into the output folder or system folders
        dirnames[:] = [d for d in dirnames
                       if d.lower() not in SKIP_DIRS and not d.startswith(".")
                       and os.path.abspath(os.path.join(dirpath, d)) != dest]
        for fn in filenames:
            if fn.startswith(".") or fn.lower() in {"desktop.ini", "thumbs.db"}:
                continue
            if fn.startswith("sort_log_") or fn.startswith("sort_plan_"):
                continue
            yield os.path.join(dirpath, fn)
        if not recursive:
            break


def unique_target(target, taken):
    base, ext = os.path.splitext(target)
    n = 1
    while os.path.exists(target) or target in taken:
        target = f"{base} ({n}){ext}"
        n += 1
    return target


def build_plan(args, categories):
    files = list(iter_files(args.source, args.recursive, args.dest))
    print(f"Found {len(files)} files. Classifying with '{args.model}'"
          + (f" + '{args.vision_model}' for images" if args.vision_model else "") + "...\n")
    plan, taken = [], set()
    for i, path in enumerate(files, 1):
        try:
            cat, reason = classify(path, args, categories)
        except Exception as e:
            cat, reason = extension_guess(path, categories), f"model error, used file type ({e})"
        target = unique_target(os.path.join(args.dest, cat, os.path.basename(path)), taken)
        taken.add(target)
        plan.append({"source": path, "target": target, "category": cat, "reason": reason})
        print(f"[{i}/{len(files)}] {os.path.basename(path)}  ->  {cat}   ({reason[:80]})")
    return plan


# ---------------------------------------------------------------- apply / undo

def apply_plan(plan, log_path):
    done = []
    for item in plan:
        if os.path.abspath(item["source"]) == os.path.abspath(item["target"]):
            continue
        os.makedirs(os.path.dirname(item["target"]), exist_ok=True)
        try:
            shutil.move(item["source"], item["target"])
            done.append({"from": item["source"], "to": item["target"]})
        except OSError as e:
            print(f"  could not move {item['source']}: {e}")
        # write the log as we go so a crash still leaves an undo-able record
        with open(log_path, "w", encoding="utf-8") as f:
            json.dump(done, f, indent=2)
    print(f"\nMoved {len(done)} files. Undo log: {log_path}")
    print(f'To undo:  python sorter.py --undo "{log_path}"')


def undo(log_path):
    with open(log_path, encoding="utf-8") as f:
        moves = json.load(f)
    restored = 0
    for m in reversed(moves):
        if not os.path.exists(m["to"]):
            print(f"  missing, skipped: {m['to']}")
            continue
        if os.path.exists(m["from"]):
            print(f"  original location occupied, skipped: {m['from']}")
            continue
        os.makedirs(os.path.dirname(m["from"]), exist_ok=True)
        shutil.move(m["to"], m["from"])
        restored += 1
    print(f"Restored {restored} of {len(moves)} files.")


# ---------------------------------------------------------------- main

def main():
    p = argparse.ArgumentParser(description="Sort files by their content using a local LLM (Ollama).")
    p.add_argument("source", nargs="?", help="folder to sort")
    p.add_argument("--dest", help="where sorted folders go (default: <source>/Sorted)")
    p.add_argument("--categories", nargs="+", default=DEFAULT_CATEGORIES,
                   help="folder names the model may choose from")
    p.add_argument("--model", default="llama3.2", help="Ollama text model (default: llama3.2)")
    p.add_argument("--vision-model", default="",
                   help="optional Ollama vision model for images, e.g. qwen2.5vl or llama3.2-vision")
    p.add_argument("--host", default="http://localhost:11434", help="Ollama server address")
    p.add_argument("--recursive", action="store_true", help="also sort files in subfolders")
    p.add_argument("--apply", action="store_true", help="actually move files (default is preview only)")
    p.add_argument("--ask", action="store_true", help="show the preview, then ask before moving")
    p.add_argument("--undo", metavar="LOG", help="undo a previous run using its sort_log_*.json")
    args = p.parse_args()

    if args.undo:
        undo(args.undo)
        return
    if not args.source or not os.path.isdir(args.source):
        p.error("please give a folder to sort")

    args.source = os.path.abspath(args.source)
    args.dest = os.path.abspath(args.dest or os.path.join(args.source, "Sorted"))
    categories = list(dict.fromkeys(args.categories))
    check_ollama(args.host, [args.model, args.vision_model])

    plan = build_plan(args, categories)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    plan_path = os.path.join(args.source, f"sort_plan_{stamp}.csv")
    with open(plan_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["source", "target", "category", "reason"])
        w.writeheader()
        w.writerows(plan)

    counts = {}
    for item in plan:
        counts[item["category"]] = counts.get(item["category"], 0) + 1
    print("\nSummary:")
    for cat, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {cat:<22} {n}")
    print(f"\nFull plan saved to {plan_path}")

    if args.ask and not args.apply:
        args.apply = input("\nMove the files as shown above? (y/n): ").strip().lower() == "y"
    if not args.apply:
        print("\nPREVIEW ONLY - nothing was moved. Re-run with --apply to move the files.")
        return
    apply_plan(plan, os.path.join(args.source, f"sort_log_{stamp}.json"))


if __name__ == "__main__":
    main()
